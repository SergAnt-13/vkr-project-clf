"""Semantic OKPD-2 classification pipeline."""

from __future__ import annotations

import logging
from dataclasses import replace
from pathlib import Path
from typing import Dict, List, Optional

import pandas as pd

from config import config
from data_loader import DataLoader
from models.reranker import SemanticReranker
from models.retriever import OKPDCandidate, SemanticRetriever
from output_manager import OutputManager
from text_processor import TextProcessor
from vat_processor import VATProcessor

logger = logging.getLogger(__name__)


class SemanticPipeline:
    """Two-stage OKPD-2 semantic search: bi-encoder retrieval + cross-encoder reranking."""

    def __init__(self, config_obj=None):
        self.config = config_obj or config

    def run(
        self,
        prediction_file: Optional[Path] = None,
        training_file: Optional[Path] = None,
        output_dir: Optional[Path] = None,
        max_rows: Optional[int] = None,
        training_max_rows: Optional[int] = None,
        limit: Optional[int] = None,
        fine_tune: bool = True,
    ) -> Dict[str, object]:
        runtime = self._runtime_config(output_dir=output_dir, training_file=training_file)
        loader = DataLoader(runtime)
        vat_processor = VATProcessor(runtime)
        output_manager = OutputManager(runtime)

        prediction_path = Path(prediction_file) if prediction_file else runtime.PRODUCTS_FILE
        prediction_raw = loader.read_table(prediction_path)
        row_limit = limit if limit is not None else max_rows
        prediction_df = loader.standardize_products(prediction_raw, max_rows=row_limit)

        training_path = Path(training_file) if training_file else runtime.MERGED_PRODUCTS_FILE
        if training_path.exists():
            training_df = loader.load_merged_products(training_path, max_rows=training_max_rows)
        else:
            logger.warning("Training file is missing: %s. Falling back to prediction data.", training_path)
            training_path = prediction_path
            training_df = loader.load_products(prediction_path, max_rows=training_max_rows)

        abbreviation_rules = loader.load_abbreviations()
        text_processor = TextProcessor(abbreviation_rules=abbreviation_rules, config_obj=runtime)
        prediction_df["name_semantic"] = text_processor.normalize_text_light_series(prediction_df["name_raw"])
        training_df["name_semantic"] = text_processor.normalize_text_light_series(training_df["name_raw"])

        reference_df = self._load_okpd_reference(loader, training_df)
        description_by_code = dict(zip(reference_df["code"], reference_df["description"]))

        retriever = SemanticRetriever(
            model_name=runtime.SEMANTIC_BI_ENCODER_MODEL,
            cache_dir=runtime.FAISS_DIR,
            finetuned_dir=runtime.SBERT_FINETUNED_DIR,
            batch_size=runtime.SEMANTIC_BATCH_SIZE,
        )
        if fine_tune:
            retriever.maybe_finetune(self._build_training_pairs(training_df, description_by_code))
        retriever.build_or_load(reference_df)

        reranker = SemanticReranker(
            model_name=runtime.SEMANTIC_CROSS_ENCODER_MODEL,
            batch_size=runtime.SEMANTIC_BATCH_SIZE,
        )

        products = prediction_df["name_semantic"].astype(str).tolist()
        retrieved = retriever.search(products, top_k=runtime.SEMANTIC_TOP_K)
        rows: List[Dict[str, object]] = []

        try:
            from tqdm import tqdm
        except ImportError:
            tqdm = lambda iterable, **_: iterable  # noqa: E731

        for product_idx, (product_text, candidates) in tqdm(
            list(enumerate(zip(products, retrieved))), desc="Semantic reranking"
        ):
            current_code = str(prediction_df.iloc[product_idx].get("okpd2_current", "")).strip()
            candidates = self._with_current_code_candidate(candidates, current_code, description_by_code)
            ranked = reranker.rerank(product_text, candidates, top_n=runtime.SEMANTIC_TOP_N_OUTPUT)
            rows.append(self._prediction_row(ranked, runtime.SEMANTIC_TOP_N_OUTPUT))
            self._clear_gpu_cache(product_idx)

        semantic_df = pd.concat([prediction_df.reset_index(drop=True), pd.DataFrame(rows)], axis=1)
        semantic_df["okpd2_pred"] = semantic_df["predicted_code"].fillna("")
        semantic_df["conf"] = semantic_df["score"].fillna(0.0)
        semantic_df["okpd2_final"] = semantic_df["okpd2_pred"]

        pp908_text = loader.load_pp908_text()
        if pp908_text:
            vat_processor.build_vat10_prefixes(pp908_text)
        semantic_df["vat_pred"] = vat_processor.process_vat_predictions(semantic_df, "okpd2_final")
        semantic_df["vat_final"] = semantic_df["vat_pred"]

        output_file = runtime.OUTPUT_DIR / f"{prediction_path.stem}_semantic.xlsx"
        self._save_semantic_excel(semantic_df, output_file)

        tables = output_manager.create_analysis_tables(semantic_df)
        saved_files = output_manager.save_tables(tables, semantic_df)
        saved_files["semantic_excel"] = output_file

        stats_for_report = {
            "classification": self._classification_stats(semantic_df),
            "vat": vat_processor.get_vat_statistics(semantic_df["vat_pred"].tolist()),
        }
        report = output_manager.generate_summary_report(
            semantic_df,
            tables,
            stats_for_report,
            metadata={
                "prediction_source": str(prediction_path),
                "training_source": str(training_path),
                "okpd_reference": str(runtime.OKPD2_REFERENCE_FILE),
            },
        )

        return {
            "data": semantic_df,
            "tables": tables,
            "report": report,
            "saved_files": saved_files,
            "output_file": output_file,
            "prediction_source": prediction_path,
            "training_source": training_path,
        }

    def _runtime_config(self, output_dir: Optional[Path], training_file: Optional[Path]):
        runtime = replace(self.config)
        base_dir = Path(output_dir) if output_dir else runtime.SEMANTIC_OUTPUT_DIR
        runtime.OUTPUT_DIR = base_dir / Path(training_file).stem if training_file else base_dir
        runtime.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        runtime.FAISS_DIR.mkdir(parents=True, exist_ok=True)
        runtime.SBERT_FINETUNED_DIR.parent.mkdir(parents=True, exist_ok=True)
        return runtime

    def _load_okpd_reference(self, loader: DataLoader, training_df: pd.DataFrame) -> pd.DataFrame:
        ref_path = self.config.OKPD2_REFERENCE_FILE
        if ref_path.exists():
            ref_df = loader.read_table(ref_path)
            columns = {str(col).strip().lower(): col for col in ref_df.columns}
            code_col = columns.get("code") or columns.get("okpd2") or columns.get("okpd")
            desc_col = columns.get("description") or columns.get("name") or columns.get("title")
            if not code_col or not desc_col:
                raise ValueError(f"OKPD-2 reference must contain code and description columns: {ref_path}")
            result = ref_df[[code_col, desc_col]].copy()
            result.columns = ["code", "description"]
            result = result.dropna(subset=["code", "description"])
            result["code"] = result["code"].astype(str).str.strip()
            result["description"] = result["description"].astype(str).str.strip()
            return result.drop_duplicates("code").reset_index(drop=True)

        logger.warning(
            "Full OKPD-2 reference is missing: %s. Building fallback descriptions from training data.",
            ref_path,
        )
        label_column = "okpd2_reference"
        if not training_df[label_column].astype(str).str.match(self.config.OKPD2_PATTERN).any():
            label_column = "okpd2_current"

        fallback = training_df.copy()
        fallback["code"] = fallback[label_column].astype(str).str.strip()
        fallback = fallback[fallback["code"].str.match(self.config.OKPD2_PATTERN)]
        fallback["name_for_description"] = fallback.get("name_semantic", fallback["name_raw"]).astype(str)
        grouped = (
            fallback.groupby("code")["name_for_description"]
            .apply(lambda values: "; ".join(list(dict.fromkeys(values.dropna().astype(str).head(5)))))
            .reset_index()
        )
        grouped.columns = ["code", "description"]
        grouped["description"] = grouped.apply(
            lambda row: f"{row['code']}: {row['description']}", axis=1
        )
        return grouped

    def _build_training_pairs(self, training_df: pd.DataFrame, description_by_code: Dict[str, str]):
        pairs = []
        for _, row in training_df.iterrows():
            reference_code = str(row.get("okpd2_reference", "")).strip()
            current_code = str(row.get("okpd2_current", "")).strip()
            code = reference_code if reference_code in description_by_code else current_code
            product = str(row.get("name_semantic") or row.get("name_raw") or "").strip()
            description = description_by_code.get(code, "")
            if product and description:
                pairs.append((product, description))
        return pairs

    def _with_current_code_candidate(
        self,
        candidates: List[OKPDCandidate],
        current_code: str,
        description_by_code: Dict[str, str],
    ) -> List[OKPDCandidate]:
        if current_code not in description_by_code:
            return candidates
        if any(candidate.code == current_code for candidate in candidates):
            return candidates
        return [
            OKPDCandidate(current_code, description_by_code[current_code], 1.0),
            *candidates,
        ]

    def _prediction_row(self, ranked: List[OKPDCandidate], top_n: int) -> Dict[str, object]:
        row: Dict[str, object] = {}
        for pos in range(top_n):
            suffix = "" if pos == 0 else f"_{pos + 1}"
            if pos < len(ranked):
                row[f"predicted_code{suffix}"] = ranked[pos].code
                row[f"score{suffix}"] = ranked[pos].retrieval_score
                row[f"description{suffix}"] = ranked[pos].description
            else:
                row[f"predicted_code{suffix}"] = ""
                row[f"score{suffix}"] = 0.0
                row[f"description{suffix}"] = ""
        row["flag"] = self._flag(float(row.get("score", 0.0)))
        return row

    def _flag(self, score: float) -> str:
        if score > self.config.SEMANTIC_CONFIDENT_THRESHOLD:
            return "Уверен"
        if score >= self.config.SEMANTIC_UNCERTAIN_THRESHOLD:
            return "Сомнительный"
        return "Неопознанный"

    def _save_semantic_excel(self, df: pd.DataFrame, output_file: Path) -> None:
        from openpyxl import load_workbook
        from openpyxl.styles import PatternFill

        output_file.parent.mkdir(parents=True, exist_ok=True)
        df.to_excel(output_file, index=False)

        workbook = load_workbook(output_file)
        worksheet = workbook.active
        header = [cell.value for cell in worksheet[1]]
        flag_col = header.index("flag") + 1 if "flag" in header else None
        if flag_col is None:
            workbook.save(output_file)
            return

        fills = {
            "Уверен": PatternFill("solid", fgColor="C6EFCE"),
            "Сомнительный": PatternFill("solid", fgColor="FFEB9C"),
            "Неопознанный": PatternFill("solid", fgColor="FFC7CE"),
        }
        for row in range(2, worksheet.max_row + 1):
            flag = worksheet.cell(row=row, column=flag_col).value
            fill = fills.get(flag)
            if fill:
                for col in range(1, worksheet.max_column + 1):
                    worksheet.cell(row=row, column=col).fill = fill
        workbook.save(output_file)

    def _classification_stats(self, df: pd.DataFrame) -> Dict[str, float]:
        total = len(df)
        valid = int(df["predicted_code"].astype(str).str.strip().ne("").sum())
        return {
            "total_predictions": total,
            "valid_predictions": valid,
            "high_confidence_predictions": int((df["score"] > self.config.SEMANTIC_CONFIDENT_THRESHOLD).sum()),
            "prediction_rate": valid / total if total else 0.0,
            "average_confidence": float(df["score"].mean()) if total else 0.0,
        }

    def _clear_gpu_cache(self, idx: int) -> None:
        if idx % 100 != 0:
            return
        try:
            import torch

            if torch.cuda.is_available():
                torch.cuda.empty_cache()
        except ImportError:
            return
